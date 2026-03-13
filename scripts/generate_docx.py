#!/usr/bin/env python3
"""
Generate a Word document for the second-order meta-analysis paper.
"""

import json
import sys
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT


def set_cell_text(cell, text, bold=False, size=9):
    """Helper to set cell text with formatting."""
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = "Times New Roman"


def add_table_from_data(doc, headers, rows, col_widths=None):
    """Add a formatted table to the document."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, size=9)
        shading = table.rows[0].cells[i]._element
        from docx.oxml.ns import qn
        from lxml import etree
        shading_elm = etree.SubElement(shading, qn("w:shd"))

    # Data rows
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            set_cell_text(table.rows[r_idx + 1].cells[c_idx], str(val), size=9)

    return table


def create_document(content_file):
    """Create the Word document from content JSON."""

    with open(content_file, "r", encoding="utf-8") as f:
        content = json.load(f)

    doc = Document()

    # --- Page setup ---
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

    # --- Styles ---
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(12)
    style.paragraph_format.line_spacing = 2.0  # Double-spaced (APA)
    style.paragraph_format.space_after = Pt(0)

    # --- Title Page ---
    for _ in range(6):
        doc.add_paragraph("")

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run(content["title"])
    run.bold = True
    run.font.size = Pt(14)
    run.font.name = "Times New Roman"

    doc.add_paragraph("")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(content.get("subtitle", ""))
    run.font.size = Pt(12)
    run.font.name = "Times New Roman"

    doc.add_paragraph("")

    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = author.add_run(content.get("author", "Taiga Capital Group — Research Division"))
    run.font.size = Pt(12)

    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date_para.add_run(content.get("date", "March 2026"))
    run.font.size = Pt(12)

    # --- Abstract ---
    doc.add_page_break()
    abs_title = doc.add_paragraph()
    abs_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = abs_title.add_run("Abstract")
    run.bold = True
    run.font.size = Pt(12)

    abs_body = doc.add_paragraph(content.get("abstract", ""))
    abs_body.paragraph_format.first_line_indent = Cm(1.27)

    if content.get("keywords"):
        kw = doc.add_paragraph()
        kw.paragraph_format.first_line_indent = Cm(1.27)
        run = kw.add_run("Keywords: ")
        run.italic = True
        run.font.size = Pt(12)
        run2 = kw.add_run(content["keywords"])
        run2.italic = True
        run2.font.size = Pt(12)

    # --- Main Body Sections ---
    for section_data in content.get("sections", []):
        level = section_data.get("level", 1)
        heading = section_data.get("heading", "")

        if level == 0:
            doc.add_page_break()

        h = doc.add_heading(heading, level=min(level, 4) if level > 0 else 1)
        for run in h.runs:
            run.font.name = "Times New Roman"
            run.font.color.rgb = RGBColor(0, 0, 0)

        for block in section_data.get("content", []):
            if isinstance(block, str):
                p = doc.add_paragraph(block)
                p.paragraph_format.first_line_indent = Cm(1.27)
            elif isinstance(block, dict):
                if block.get("type") == "table":
                    add_table_from_data(
                        doc,
                        block["headers"],
                        block["rows"],
                    )
                    if block.get("note"):
                        note_p = doc.add_paragraph()
                        run = note_p.add_run(f"Note. {block['note']}")
                        run.italic = True
                        run.font.size = Pt(10)
                elif block.get("type") == "figure_placeholder":
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run(f"[{block.get('caption', 'Figure')}]")
                    run.italic = True

    # --- References ---
    doc.add_page_break()
    ref_h = doc.add_heading("References", level=1)
    for run in ref_h.runs:
        run.font.name = "Times New Roman"
        run.font.color.rgb = RGBColor(0, 0, 0)

    for ref in content.get("references", []):
        p = doc.add_paragraph(ref)
        p.paragraph_format.first_line_indent = Cm(-1.27)  # hanging indent
        p.paragraph_format.left_indent = Cm(1.27)
        for run in p.runs:
            run.font.size = Pt(12)

    # Save
    output_path = content.get("output_path", "/home/user/taiya22/docs/meta_analysis.docx")
    doc.save(output_path)
    print(f"Document saved to {output_path}")
    return output_path


if __name__ == "__main__":
    if len(sys.argv) > 1:
        create_document(sys.argv[1])
    else:
        print("Usage: python generate_docx.py <content.json>")
